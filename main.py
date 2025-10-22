# main.py
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import threading
import os
import re

# Importar lógica
from analisis import AnalizadorSentimientos as AnalizadorLogica, verificar_dependencias

# IMPORTAR EL MÓDULO DE LIMPIEZA ROBUSTA
from limpieza_robusta import LimpiadorDatosRobusto

# IMPORTAR EL MÓDULO DE MÉTRICAS
from metricas_evaluacion import CalculadorMetricas, calcular_metricas_modelo

# Dependencias opcionales
DOCX_DISPONIBLE = False
PDF_DISPONIBLE = False
SPELLCHECKER_DISPONIBLE = False

try:
    import docx
    DOCX_DISPONIBLE = True
except ImportError:
    pass

try:
    import PyPDF2
    PDF_DISPONIBLE = True
except ImportError:
    pass

# --- NUEVOS MODELOS DE ANÁLISIS DE SENTIMIENTOS ---
try:
    from pysentimiento import create_analyzer
    PYSENTIMIENTOS_DISPONIBLE = True
except ImportError:
    PYSENTIMIENTOS_DISPONIBLE = False

try:
    from transformers import pipeline
    TRANSFORMERS_DISPONIBLE = True
except ImportError:
    TRANSFORMERS_DISPONIBLE = False

# NUEVOS: VADER Y NAIVE BAYES
try:
    from nltk.sentiment.vader import SentimentIntensityAnalyzer
    VADER_DISPONIBLE = True
except ImportError:
    VADER_DISPONIBLE = False

try:
    from sklearn.feature_extraction.text import CountVectorizer
    from sklearn.naive_bayes import BernoulliNB
    from sklearn.model_selection import train_test_split
    from sklearn.metrics import accuracy_score, classification_report
    from sklearn.pipeline import Pipeline
    SKLEARN_DISPONIBLE = True
except ImportError:
    SKLEARN_DISPONIBLE = False

try:
    from spellchecker import SpellChecker
    SPELLCHECKER_DISPONIBLE = True
except ImportError:
    pass


class AnalizadorSentimientosGUI:
    def __init__(self):
        self.logica = AnalizadorLogica()
        # INSTANCIAR EL LIMPIADOR ROBUSTO DESDE EL MÓDULO IMPORTADO
        self.limpiador_robusto = LimpiadorDatosRobusto()
        # INSTANCIAR EL CALCULADOR DE MÉTRICAS
        self.calculador_metricas = CalculadorMetricas()
        self.correcciones_aplicadas = False
        self.analisis_completado = False
        self.limpieza_robusta_aplicada = False
        self.metricas_calculadas = False
        self.setup_gui()

    def setup_gui(self):
        self.ventana = tk.Tk()
        self.ventana.title("✨ Analizador de Sentimientos Profesional ")
        self.ventana.geometry("1350x950")
        self.ventana.minsize(1200, 800)  # Tamaño mínimo para evitar problemas
        
        # Esquema de colores profesional - Clean White Theme
        self.colores = {
            'bg_primary': '#ffffff',      # Fondo principal blanco
            'bg_secondary': '#f8fafc',    # Fondo secundario gris muy claro
            'bg_card': '#ffffff',         # Tarjetas blancas
            'bg_input': '#f1f5f9',        # Campos de entrada
            'bg_hover': '#f1f5f9',        # Hover suave
            'text_primary': '#0f172a',    # Texto principal negro
            'text_secondary': '#475569',  # Texto secundario gris
            'text_muted': '#64748b',      # Texto atenuado
            'accent_blue': '#2563eb',     # Azul profesional
            'accent_blue_hover': '#1d4ed8', # Azul hover
            'accent_green': '#059669',    # Verde éxito
            'accent_orange': '#d97706',   # Naranja advertencia
            'accent_red': '#dc2626',      # Rojo error
            'accent_purple': '#7c3aed',   # Púrpura para limpieza
            'accent_purple_hover': '#6d28d9', # Púrpura hover
            'border': '#e2e8f0',          # Bordes suaves
            'border_focus': '#3b82f6',    # Borde en foco
            'shadow': 'rgba(0,0,0,0.1)'   # Sombras suaves
        }
        self.ventana.configure(bg=self.colores['bg_primary'])

        # Estilos personalizados
        style = ttk.Style()
        style.theme_use('default')
        self.configurar_estilos_ttk(style)

        # === CONFIGURACIÓN MEJORADA DEL SCROLL ===
        self.setup_scroll_container()
        
        # === Construir toda la interfaz ===
        self.crear_interfaz_completa()

    def setup_scroll_container(self):
        """Configura el contenedor de scroll mejorado"""
        # Frame principal que contiene todo
        self.main_container = tk.Frame(self.ventana, bg=self.colores['bg_primary'])
        self.main_container.pack(fill='both', expand=True)
        
        # Canvas para el scroll
        self.canvas = tk.Canvas(
            self.main_container, 
            bg=self.colores['bg_primary'], 
            highlightthickness=0,
            relief='flat',
            bd=0
        )
        
        # Scrollbar vertical mejorada
        self.v_scrollbar = tk.Scrollbar(
            self.main_container, 
            orient="vertical", 
            command=self.canvas.yview,
            bg=self.colores['bg_secondary'],
            troughcolor=self.colores['bg_input'],
            activebackground=self.colores['accent_blue'],
            width=16
        )
        
        # Frame que contendrá todo el contenido
        self.scrollable_frame = tk.Frame(self.canvas, bg=self.colores['bg_primary'])
        
        # Configurar el frame dentro del canvas
        self.canvas_window = self.canvas.create_window(
            (0, 0), 
            window=self.scrollable_frame, 
            anchor="nw"
        )
        
        # Configurar scroll
        self.canvas.configure(yscrollcommand=self.v_scrollbar.set)
        
        # Empaquetar elementos
        self.v_scrollbar.pack(side="right", fill="y")
        self.canvas.pack(side="left", fill="both", expand=True)
        
        # Configurar eventos de scroll
        self.configurar_eventos_scroll()
        
        # Configurar actualización automática del scroll
        self.configurar_actualizacion_scroll()

    def configurar_eventos_scroll(self):
        """Configura los eventos de scroll del mouse"""
        def _on_mousewheel(event):
            # Verificar si el canvas puede hacer scroll
            if self.canvas.bbox("all"):
                self.canvas.yview_scroll(int(-1*(event.delta/120)), "units")
        
        def _on_mousewheel_linux_up(event):
            if self.canvas.bbox("all"):
                self.canvas.yview_scroll(-1, "units")
                
        def _on_mousewheel_linux_down(event):
            if self.canvas.bbox("all"):
                self.canvas.yview_scroll(1, "units")
        
        # Bind a todo el canvas y frame
        self.canvas.bind_all("<MouseWheel>", _on_mousewheel)  # Windows
        self.canvas.bind_all("<Button-4>", _on_mousewheel_linux_up)  # Linux up
        self.canvas.bind_all("<Button-5>", _on_mousewheel_linux_down)  # Linux down
        
        # También bind al frame scrollable
        self.scrollable_frame.bind_all("<MouseWheel>", _on_mousewheel)
        self.scrollable_frame.bind_all("<Button-4>", _on_mousewheel_linux_up)
        self.scrollable_frame.bind_all("<Button-5>", _on_mousewheel_linux_down)

    def configurar_actualizacion_scroll(self):
        """Configura la actualización automática del scroll"""
        def configure_scroll(event=None):
            # Actualizar el tamaño del canvas window
            canvas_width = self.canvas.winfo_width()
            self.canvas.itemconfig(self.canvas_window, width=canvas_width)
            
            # Actualizar scroll region
            self.canvas.configure(scrollregion=self.canvas.bbox("all"))
        
        # Bind para cuando cambie el tamaño del canvas
        self.canvas.bind('<Configure>', configure_scroll)
        
        # Bind para cuando cambie el contenido del frame
        def update_scroll_region(event=None):
            self.canvas.update_idletasks()
            self.canvas.configure(scrollregion=self.canvas.bbox("all"))
        
        self.scrollable_frame.bind('<Configure>', update_scroll_region)
        
        # Actualización inicial después de que se cree todo
        def initial_update():
            self.canvas.update_idletasks()
            configure_scroll()
            update_scroll_region()
        
        self.ventana.after(100, initial_update)

    def crear_interfaz_completa(self):
        """Crea toda la interfaz dentro del frame scrollable"""
        # Container principal con padding
        main_content = tk.Frame(self.scrollable_frame, bg=self.colores['bg_primary'])
        main_content.pack(fill='both', expand=True, padx=40, pady=30)
        
        # Header profesional mejorado
        self.crear_header_profesional(main_content)
        
        # Separador elegante
        self.crear_separador(main_content)
        
        # Panel de controles profesional ampliado
        self.crear_panel_controles_ampliado(main_content)
        
        # Panel de información dual mejorado
        self.crear_panel_informacion_mejorado(main_content)
        
        # Área de resultados premium
        self.crear_area_resultados_premium(main_content)
        
        # Footer con progreso profesional
        self.crear_footer_profesional(main_content)
        
        # Actualizar scroll después de crear todo
        self.ventana.after(200, lambda: self.canvas.configure(scrollregion=self.canvas.bbox("all")))

    def configurar_estilos_ttk(self, style):
        """Configura estilos TTK profesionales"""
        # Notebook elegante
        style.configure('Professional.TNotebook', 
                       background=self.colores['bg_primary'], 
                       borderwidth=0,
                       tabmargins=[2, 5, 2, 0])
        style.configure('Professional.TNotebook.Tab',
                       background=self.colores['bg_input'],
                       foreground=self.colores['text_secondary'],
                       padding=[25, 12],
                       focuscolor='none',
                       font=('Segoe UI', 10, 'normal'))
        style.map('Professional.TNotebook.Tab',
                 background=[('selected', self.colores['accent_blue']),
                           ('active', self.colores['bg_hover'])],
                 foreground=[('selected', 'white'),
                           ('active', self.colores['text_primary'])])
        
        # Progressbar profesional
        style.configure('Professional.Horizontal.TProgressbar',
                       background=self.colores['accent_blue'],
                       troughcolor=self.colores['bg_input'],
                       borderwidth=0,
                       lightcolor=self.colores['accent_blue'],
                       darkcolor=self.colores['accent_blue'],
                       relief='flat')

    def crear_header_profesional(self, parent):
        """Crea un header profesional y elegante mejorado"""
        header_container = tk.Frame(parent, bg=self.colores['bg_primary'])
        header_container.pack(fill='x', pady=(0, 25))
        
        # Logo y título
        title_frame = tk.Frame(header_container, bg=self.colores['bg_primary'])
        title_frame.pack(fill='x')
        
        # Título principal con gradiente visual
        title_container = tk.Frame(title_frame, bg=self.colores['bg_primary'])
        title_container.pack(anchor='w')
        
        main_title = tk.Label(
            title_container,
            text="✨ Analizador de Sentimientos",
            font=('Segoe UI', 32, 'bold'),
            bg=self.colores['bg_primary'],
            fg=self.colores['text_primary']
        )
        main_title.pack(side='left')
        
        # Subtítulo profesional mejorado
        subtitle = tk.Label(
            title_frame,
            text="Análisis Inteligente de Emociones • Procesamiento Avanzado de Lenguaje Natural • Corrección Ortográfica Automática • Evaluación de Métricas",
            font=('Segoe UI', 12),
            bg=self.colores['bg_primary'],
            fg=self.colores['text_secondary']
        )
        subtitle.pack(anchor='w', pady=(8, 0))
        
        # Badge de funcionalidades
        badge_container = tk.Frame(title_frame, bg=self.colores['bg_primary'])
        badge_container.pack(anchor='w', pady=(5, 0))
        
        if SPELLCHECKER_DISPONIBLE:
            spell_badge = tk.Label(
                badge_container,
                text="✏️ Corrector ortográfico",
                font=('Segoe UI', 9, 'bold'),
                bg=self.colores['accent_green'],
                fg='white',
                padx=10,
                pady=3
            )
            spell_badge.pack(side='left', padx=(0, 5))
        
        clean_badge = tk.Label(
            badge_container,
            text="🧹 Limpiador Robusto",
            font=('Segoe UI', 9, 'bold'),
            bg=self.colores['accent_purple'],
            fg='white',
            padx=10,
            pady=3
        )
        clean_badge.pack(side='left', padx=(0, 5))
        
        # NUEVOS BADGES PARA VADER Y NAIVE BAYES
        if VADER_DISPONIBLE:
            vader_badge = tk.Label(
                badge_container,
                text="⚡ VADER",
                font=('Segoe UI', 9, 'bold'),
                bg='#8B4513',
                fg='white',
                padx=10,
                pady=3
            )
            vader_badge.pack(side='left', padx=(0, 5))
        
        if SKLEARN_DISPONIBLE:
            nb_badge = tk.Label(
                badge_container,
                text="📊 Naive Bayes",
                font=('Segoe UI', 9, 'bold'),
                bg='#2F4F4F',
                fg='white',
                padx=10,
                pady=3
            )
            nb_badge.pack(side='left', padx=(0, 5))
        
        # Badge para métricas
        metrics_badge = tk.Label(
            badge_container,
            text="📈 Métricas Avanzadas",
            font=('Segoe UI', 9, 'bold'),
            bg='#7C3AED',
            fg='white',
            padx=10,
            pady=3
        )
        metrics_badge.pack(side='left', padx=(0, 5))

    def crear_separador(self, parent):
        """Crea un separador elegante"""
        separator = tk.Frame(parent, height=2, bg=self.colores['border'])
        separator.pack(fill='x', pady=(0, 30))

    def crear_panel_controles_ampliado(self, parent):
        """Crea panel de controles profesional ampliado con limpieza robusta post-análisis"""
        controls_container = tk.Frame(parent, bg=self.colores['bg_primary'])
        controls_container.pack(fill='x', pady=(0, 30))
        
        # Título de sección
        section_title = tk.Label(
            controls_container,
            text="🎛️ Panel de Control Avanzado",
            font=('Segoe UI', 16, 'bold'),
            bg=self.colores['bg_primary'],
            fg=self.colores['text_primary']
        )
        section_title.pack(anchor='w', pady=(0, 15))
        
        # Container de botones con grid ampliado (3 filas)
        buttons_grid = tk.Frame(controls_container, bg=self.colores['bg_primary'])
        buttons_grid.pack(fill='x')
        
        # Configurar grid para 3 filas y 3 columnas
        for i in range(3):
            buttons_grid.columnconfigure(i, weight=1)
        
        # Primera fila de botones
        self.crear_boton_profesional(buttons_grid, "📁 Cargar Archivo", 
                                    "Importar datos para análisis", 
                                    self.cargar_archivo, 0, 0)
        
        self.crear_boton_profesional(buttons_grid, "🔍 Analizar", 
                                    "Procesar sentimientos", 
                                    self.analizar_sentimientos, 1, 0, 
                                    state='disabled')
        
        self.crear_boton_profesional(buttons_grid, "🧹 Limpiar Datos", 
                                    "Corrección ortográfica y limpieza", 
                                    self.limpiar_datos_robusto, 2, 0, 
                                    state='disabled', color='purple')
        
        # Segunda fila de botones
        self.crear_boton_profesional(buttons_grid, "📊 Métricas", 
                                    "Evaluar rendimiento del modelo", 
                                    self.calcular_metricas_evaluacion, 0, 1, 
                                    state='disabled', color='blue')
        
        self.crear_boton_profesional(buttons_grid, "📈 Visualizar", 
                                    "Mostrar gráficos avanzados", 
                                    self.mostrar_graficos, 1, 1, 
                                    state='disabled', color='orange')
        
        self.crear_boton_profesional(buttons_grid, "💾 Exportar", 
                                    "Guardar resultados completos", 
                                    self.exportar_resultados, 2, 1, 
                                    state='disabled', color='green')
        
        # Tercera fila de botones
        self.crear_boton_profesional(buttons_grid, "🔄 Resetear", 
                                    "Limpiar y comenzar de nuevo", 
                                    self.resetear_analisis, 0, 2, 
                                    state='disabled', color='red')

    def crear_boton_profesional(self, parent, titulo, descripcion, comando, col, row, color='blue', **kwargs):
        """Crea un botón con diseño profesional y descripción con colores personalizables"""
        # Container principal
        btn_container = tk.Frame(parent, bg=self.colores['bg_primary'])
        btn_container.grid(row=row, column=col, padx=12, pady=10, sticky="ew")
        
        # Card del botón
        card = tk.Frame(
            btn_container,
            bg=self.colores['bg_card'],
            relief='solid',
            bd=1,
            highlightbackground=self.colores['border'],
            highlightthickness=1
        )
        card.pack(fill='both', expand=True)
        
        # Seleccionar colores según el tipo
        color_map = {
            'blue': (self.colores['accent_blue'], self.colores['accent_blue_hover']),
            'purple': (self.colores['accent_purple'], self.colores['accent_purple_hover']),
            'green': (self.colores['accent_green'], '#047857'),
            'orange': (self.colores['accent_orange'], '#b45309'),
            'red': (self.colores['accent_red'], '#b91c1c')
        }
        bg_color, hover_color = color_map.get(color, color_map['blue'])
        
        # Botón principal
        main_btn = tk.Button(
            card,
            text=titulo,
            command=comando,
            font=('Segoe UI', 11, 'bold'),
            bg=bg_color,
            fg='white',
            relief='flat',
            bd=0,
            padx=20,
            pady=15,
            cursor='hand2',
            activebackground=hover_color,
            activeforeground='white',
            **kwargs
        )
        main_btn.pack(fill='x', padx=15, pady=(15, 5))
        
        # Descripción
        desc_label = tk.Label(
            card,
            text=descripcion,
            font=('Segoe UI', 9),
            bg=self.colores['bg_card'],
            fg=self.colores['text_muted'],
            wraplength=180
        )
        desc_label.pack(pady=(0, 15))
        
        # Efectos hover para la card completa
        def on_enter_card(e):
            if main_btn['state'] != 'disabled':
                card.config(highlightbackground=self.colores['border_focus'])
                main_btn.config(bg=hover_color)
        
        def on_leave_card(e):
            if main_btn['state'] != 'disabled':
                card.config(highlightbackground=self.colores['border'])
                main_btn.config(bg=bg_color)
        
        card.bind("<Enter>", on_enter_card)
        card.bind("<Leave>", on_leave_card)
        main_btn.bind("<Enter>", on_enter_card)
        main_btn.bind("<Leave>", on_leave_card)
        
        # Guardar referencias de botones
        if 'Cargar' in titulo:
            self.btn_cargar = main_btn
        elif 'Analizar' in titulo:
            self.btn_analizar = main_btn
        elif 'Limpiar' in titulo:
            self.btn_limpiar = main_btn
        elif 'Métricas' in titulo:
            self.btn_metricas = main_btn
        elif 'Exportar' in titulo:
            self.btn_exportar = main_btn
        elif 'Visualizar' in titulo:
            self.btn_graficos = main_btn
        elif 'Resetear' in titulo:
            self.btn_resetear = main_btn

    def crear_panel_informacion_mejorado(self, parent):
        """Crea panel de información profesional mejorado - MANTIENE EL DASHBOARD ORIGINAL"""
        info_container = tk.Frame(parent, bg=self.colores['bg_primary'])
        info_container.pack(fill='x', pady=(0, 30))
        
        # Título de sección
        section_title = tk.Label(
            info_container,
            text="📊 Dashboard de Información del Proyecto",
            font=('Segoe UI', 16, 'bold'),
            bg=self.colores['bg_primary'],
            fg=self.colores['text_primary']
        )
        section_title.pack(anchor='w', pady=(0, 15))
        
        # Grid de información (3 columnas) - MANTIENE LAS 3 COLUMNAS ORIGINALES
        info_grid = tk.Frame(info_container, bg=self.colores['bg_primary'])
        info_grid.pack(fill='x')
        info_grid.columnconfigure(0, weight=1)
        info_grid.columnconfigure(1, weight=1)
        info_grid.columnconfigure(2, weight=1)
        
        # Card de información del archivo (ORIGINAL)
        self.crear_card_informacion_mejorada(info_grid, 0)
        
        # Card de estadísticas (ORIGINAL - MÉTRICAS AVANZADAS)
        self.crear_card_estadisticas_mejorada(info_grid, 1)
        
        # Card de corrección ortográfica (ORIGINAL)
        self.crear_card_correccion_ortografica(info_grid, 2)

    def crear_card_informacion_mejorada(self, parent, columna):
        """Crea card de información del archivo mejorada - MANTIENE ORIGINAL"""
        card = self.crear_card_base(parent, columna)
        
        # Header de la card con icono animado
        header = tk.Frame(card, bg=self.colores['bg_card'])
        header.pack(fill='x', pady=(20, 10), padx=25)
        
        icon_title = tk.Frame(header, bg=self.colores['bg_card'])
        icon_title.pack(fill='x')
        
        # Icono con color dinámico
        self.icon_archivo = tk.Label(
            icon_title,
            text="📄",
            font=('Segoe UI', 20),
            bg=self.colores['bg_card'],
            fg=self.colores['text_secondary']
        )
        self.icon_archivo.pack(side='left', padx=(0, 10))
        
        tk.Label(
            icon_title,
            text="Archivo de Datos",
            font=('Segoe UI', 14, 'bold'),
            bg=self.colores['bg_card'],
            fg=self.colores['text_primary']
        ).pack(side='left', anchor='w')
        
        # Badge de estado
        self.status_badge = tk.Label(
            icon_title,
            text="Esperando",
            font=('Segoe UI', 8, 'bold'),
            bg=self.colores['text_muted'],
            fg='white',
            padx=8,
            pady=2
        )
        self.status_badge.pack(side='right')
        
        # Separador interno
        sep = tk.Frame(card, height=1, bg=self.colores['border'])
        sep.pack(fill='x', padx=25, pady=(0, 15))
        
        # Contenido - MANTIENE TEXTO ORIGINAL
        content = tk.Frame(card, bg=self.colores['bg_card'])
        content.pack(fill='both', expand=True, padx=25, pady=(0, 20))
        
        self.info_archivo = tk.Label(
            content,
            text="📁 No hay archivo seleccionado\n• Selecciona un archivo de texto, CSV, Excel o Word\n• Formatos soportados: TXT, CSV, XLSX, DOCX, PDF\n• Tamaño máximo recomendado: 50MB\n• Corrección ortográfica automática disponible",
            font=('Segoe UI', 11),
            bg=self.colores['bg_card'],
            fg=self.colores['text_secondary'],
            justify='left',
            wraplength=320
        )
        self.info_archivo.pack(anchor='w', fill='both')

    def crear_card_estadisticas_mejorada(self, parent, columna):
        """Crea card de estadísticas mejorada - MANTIENE ORIGINAL"""
        card = self.crear_card_base(parent, columna)
        
        # Header de la card
        header = tk.Frame(card, bg=self.colores['bg_card'])
        header.pack(fill='x', pady=(20, 10), padx=25)
        
        icon_title = tk.Frame(header, bg=self.colores['bg_card'])
        icon_title.pack(fill='x')
        
        tk.Label(
            icon_title,
            text="📈",
            font=('Segoe UI', 20),
            bg=self.colores['bg_card'],
            fg=self.colores['accent_green']
        ).pack(side='left', padx=(0, 10))
        
        tk.Label(
            icon_title,
            text="Métricas Avanzadas",
            font=('Segoe UI', 14, 'bold'),
            bg=self.colores['bg_card'],
            fg=self.colores['text_primary']
        ).pack(side='left', anchor='w')
        
        # Badge de análisis
        self.analysis_badge = tk.Label(
            icon_title,
            text="Pendiente",
            font=('Segoe UI', 8, 'bold'),
            bg=self.colores['text_muted'],
            fg='white',
            padx=8,
            pady=2
        )
        self.analysis_badge.pack(side='right')
        
        # Separador interno
        sep = tk.Frame(card, height=1, bg=self.colores['border'])
        sep.pack(fill='x', padx=25, pady=(0, 15))
        
        # Contenido - MANTIENE TEXTO ORIGINAL
        content = tk.Frame(card, bg=self.colores['bg_card'])
        content.pack(fill='both', expand=True, padx=25, pady=(0, 20))
        
        self.stats_label = tk.Label(
            content,
            text="📊 Ejecuta el análisis para ver estadísticas\n• Distribución de sentimientos\n• Puntuaciones de confianza\n• Métricas de intensidad emocional\n• Análisis de correlaciones\n• Insights detallados",
            font=('Segoe UI', 11),
            bg=self.colores['bg_card'],
            fg=self.colores['text_secondary'],
            justify='left',
            wraplength=320
        )
        self.stats_label.pack(anchor='w', fill='both')

    def crear_card_correccion_ortografica(self, parent, columna):
        """Crea card de corrección ortográfica - MANTIENE ORIGINAL"""
        card = self.crear_card_base(parent, columna)
        
        # Header de la card
        header = tk.Frame(card, bg=self.colores['bg_card'])
        header.pack(fill='x', pady=(20, 10), padx=25)
        
        icon_title = tk.Frame(header, bg=self.colores['bg_card'])
        icon_title.pack(fill='x')
        
        self.icon_limpieza = tk.Label(
            icon_title,
            text="✏️",
            font=('Segoe UI', 20),
            bg=self.colores['bg_card'],
            fg=self.colores['accent_purple']
        )
        self.icon_limpieza.pack(side='left', padx=(0, 10))
        
        tk.Label(
            icon_title,
            text="Corrección Ortográfica",
            font=('Segoe UI', 14, 'bold'),
            bg=self.colores['bg_card'],
            fg=self.colores['text_primary']
        ).pack(side='left', anchor='w')
        
        # Badge de limpieza
        self.cleaning_badge = tk.Label(
            icon_title,
            text="Pendiente",
            font=('Segoe UI', 8, 'bold'),
            bg=self.colores['text_muted'],
            fg='white',
            padx=8,
            pady=2
        )
        self.cleaning_badge.pack(side='right')
        
        # Separador interno
        sep = tk.Frame(card, height=1, bg=self.colores['border'])
        sep.pack(fill='x', padx=25, pady=(0, 15))
        
        # Contenido - MANTIENE TEXTO ORIGINAL
        content = tk.Frame(card, bg=self.colores['bg_card'])
        content.pack(fill='both', expand=True, padx=25, pady=(0, 20))
        
        texto_limpieza = "✏️ Corrección automática de ortografía\n• Disponible tras completar análisis\n• 'me seto felis' → 'me siento feliz'\n• Conversión de emojis a texto\n• Limpieza de URLs y emails\n• Preservación del contexto emocional"
        
        self.limpieza_label = tk.Label(
            content,
            text=texto_limpieza,
            font=('Segoe UI', 11),
            bg=self.colores['bg_card'],
            fg=self.colores['text_secondary'],
            justify='left',
            wraplength=320
        )
        self.limpieza_label.pack(anchor='w', fill='both')

    def crear_card_base(self, parent, columna):
        """Crea una card base profesional"""
        card = tk.Frame(
            parent,
            bg=self.colores['bg_card'],
            relief='solid',
            bd=1,
            highlightbackground=self.colores['border'],
            highlightthickness=1
        )
        card.grid(row=0, column=columna, padx=10, pady=0, sticky="nsew")
        return card

    def crear_area_resultados_premium(self, parent):
        """Crea área de resultados premium con scroll mejorado"""
        results_container = tk.Frame(parent, bg=self.colores['bg_primary'])
        results_container.pack(fill='both', expand=True, pady=(0, 25))
        
        # Título de sección
        section_title = tk.Label(
            results_container,
            text="📋 Resultados del Análisis Avanzado",
            font=('Segoe UI', 16, 'bold'),
            bg=self.colores['bg_primary'],
            fg=self.colores['text_primary']
        )
        section_title.pack(anchor='w', pady=(0, 15))
        
        # Container de pestañas con altura fija para mejor scroll
        tabs_container = tk.Frame(
            results_container,
            bg=self.colores['bg_secondary'],
            relief='solid',
            bd=1,
            highlightbackground=self.colores['border'],
            highlightthickness=1
        )
        tabs_container.pack(fill='both', expand=True)
        
        # Configurar altura mínima
        tabs_container.configure(height=500)
        
        # Notebook profesional
        self.notebook = ttk.Notebook(tabs_container, style='Professional.TNotebook')
        self.notebook.pack(fill='both', expand=True, padx=20, pady=20)
        
        # Crear pestañas premium mejoradas
        self.crear_pestana_resumen_premium()
        self.crear_pestana_datos_premium()
        self.crear_pestana_limpieza_robusta()
        # NUEVA PESTAÑA PARA MÉTRICAS
        self.crear_pestana_metricas_evaluacion()

    def crear_pestana_resumen_premium(self):
        """Crea pestaña de resumen premium con mejor scroll"""
        resumen_frame = tk.Frame(self.notebook, bg=self.colores['bg_primary'])
        self.notebook.add(resumen_frame, text="📋 Resumen Ejecutivo")
        
        # Container principal con configuración mejorada
        main_container = tk.Frame(resumen_frame, bg=self.colores['bg_primary'])
        main_container.pack(fill='both', expand=True, padx=15, pady=15)
        
        # Frame para el texto y scrollbar
        text_frame = tk.Frame(main_container, bg=self.colores['bg_primary'])
        text_frame.pack(fill='both', expand=True)
        
        # Área de texto premium con mejor configuración
        self.texto_resumen = tk.Text(
            text_frame,
            wrap='word',
            font=('Segoe UI', 11),
            bg=self.colores['bg_primary'],
            fg=self.colores['text_primary'],
            relief='flat',
            bd=0,
            padx=15,
            pady=15,
            spacing1=4,
            spacing2=2,
            spacing3=1,
            insertbackground=self.colores['text_primary'],
            selectbackground=self.colores['accent_blue'],
            selectforeground='white',
            height=20
        )
        self.texto_resumen.pack(side='left', fill='both', expand=True)
        
        # Scrollbar más visible
        scrollbar_resumen = tk.Scrollbar(
            text_frame,
            orient='vertical',
            command=self.texto_resumen.yview,
            bg=self.colores['bg_secondary'],
            troughcolor=self.colores['bg_input'],
            activebackground=self.colores['accent_blue'],
            width=14
        )
        scrollbar_resumen.pack(side='right', fill='y', padx=(5, 0))
        self.texto_resumen.config(yscrollcommand=scrollbar_resumen.set)
        
        # Texto inicial
        self.mostrar_en_resumen("📋 RESUMEN EJECUTIVO\nAquí se mostrará un resumen completo del análisis una vez que cargues un archivo y ejecutes el análisis de sentimientos con corrección ortográfica automática.")

    def crear_pestana_datos_premium(self):
        """Crea pestaña de datos premium con mejor scroll"""
        datos_frame = tk.Frame(self.notebook, bg=self.colores['bg_primary'])
        self.notebook.add(datos_frame, text="📊 Datos Detallados")
        
        # Container principal
        main_container = tk.Frame(datos_frame, bg=self.colores['bg_primary'])
        main_container.pack(fill='both', expand=True, padx=15, pady=15)
        
        # Frame para el texto y scrollbars
        text_frame = tk.Frame(main_container, bg=self.colores['bg_primary'])
        text_frame.pack(fill='both', expand=True)
        
        # Área de texto premium con scrolling horizontal y vertical
        self.texto_datos = tk.Text(
            text_frame,
            wrap='none',
            font=('Consolas', 10),
            bg=self.colores['bg_primary'],
            fg=self.colores['text_primary'],
            relief='flat',
            bd=0,
            padx=15,
            pady=15,
            insertbackground=self.colores['text_primary'],
            selectbackground=self.colores['accent_blue'],
            selectforeground='white',
            height=20
        )
        self.texto_datos.grid(row=0, column=0, sticky='nsew')
        
        # Scrollbars
        v_scrollbar = tk.Scrollbar(
            text_frame,
            orient='vertical',
            command=self.texto_datos.yview,
            bg=self.colores['bg_secondary'],
            troughcolor=self.colores['bg_input'],
            activebackground=self.colores['accent_blue'],
            width=14
        )
        v_scrollbar.grid(row=0, column=1, sticky='ns')
        self.texto_datos.config(yscrollcommand=v_scrollbar.set)
        
        h_scrollbar = tk.Scrollbar(
            text_frame,
            orient='horizontal',
            command=self.texto_datos.xview,
            bg=self.colores['bg_secondary'],
            troughcolor=self.colores['bg_input'],
            activebackground=self.colores['accent_blue']
        )
        h_scrollbar.grid(row=1, column=0, sticky='ew')
        self.texto_datos.config(xscrollcommand=h_scrollbar.set)
        
        # Configurar expansión del grid
        text_frame.rowconfigure(0, weight=1)
        text_frame.columnconfigure(0, weight=1)
        
        # Texto inicial
        self.mostrar_en_datos("📊 DATOS DETALLADOS\nAquí se mostrarán los datos detallados del análisis con todas las métricas calculadas y correcciones ortográficas aplicadas.")

    def crear_pestana_limpieza_robusta(self):
        """Crea pestaña específica para mostrar limpieza robusta"""
        limpieza_frame = tk.Frame(self.notebook, bg=self.colores['bg_primary'])
        self.notebook.add(limpieza_frame, text="✏️ Corrección Ortográfica")
        
        # Container principal
        main_container = tk.Frame(limpieza_frame, bg=self.colores['bg_primary'])
        main_container.pack(fill='both', expand=True, padx=15, pady=15)
        
        # Frame para el texto y scrollbar
        text_frame = tk.Frame(main_container, bg=self.colores['bg_primary'])
        text_frame.pack(fill='both', expand=True)
        
        # Área de texto para limpieza robusta
        self.texto_limpieza_robusta = tk.Text(
            text_frame,
            wrap='word',
            font=('Segoe UI', 11),
            bg=self.colores['bg_primary'],
            fg=self.colores['text_primary'],
            relief='flat',
            bd=0,
            padx=15,
            pady=15,
            spacing1=4,
            spacing2=2,
            spacing3=1,
            insertbackground=self.colores['text_primary'],
            selectbackground=self.colores['accent_purple'],
            selectforeground='white',
            height=20
        )
        self.texto_limpieza_robusta.pack(side='left', fill='both', expand=True)
        
        # Scrollbar
        scrollbar_limpieza = tk.Scrollbar(
            text_frame,
            orient='vertical',
            command=self.texto_limpieza_robusta.yview,
            bg=self.colores['bg_secondary'],
            troughcolor=self.colores['bg_input'],
            activebackground=self.colores['accent_purple'],
            width=14
        )
        scrollbar_limpieza.pack(side='right', fill='y', padx=(5, 0))
        self.texto_limpieza_robusta.config(yscrollcommand=scrollbar_limpieza.set)
        
        # Texto inicial
        self.mostrar_en_limpieza_robusta("✏️ CORRECCIÓN ORTOGRÁFICA AUTOMÁTICA\n" + 
                                         "Esta función estará disponible después de completar el análisis de sentimientos.\n" +
                                         "📝 La corrección ortográfica incluye:\n" +
                                         "• Corrección de palabras mal escritas: 'me seto felis' → 'me siento feliz'\n" +
                                         "• Normalización de texto: 'estoi' → 'estoy'\n" +
                                         "• Conversión de emojis a texto descriptivo\n" +
                                         "• Eliminación de URLs y enlaces\n" +
                                         "• Limpieza de direcciones de email\n" +
                                         "• Normalización de caracteres especiales\n" +
                                         "• Preservación del contexto emocional")

    def crear_pestana_metricas_evaluacion(self):
        """Crea pestaña específica para mostrar métricas de evaluación"""
        metricas_frame = tk.Frame(self.notebook, bg=self.colores['bg_primary'])
        self.notebook.add(metricas_frame, text="📈 Métricas de Evaluación")
        
        # Container principal
        main_container = tk.Frame(metricas_frame, bg=self.colores['bg_primary'])
        main_container.pack(fill='both', expand=True, padx=15, pady=15)
        
        # Frame para el texto y scrollbar
        text_frame = tk.Frame(main_container, bg=self.colores['bg_primary'])
        text_frame.pack(fill='both', expand=True)
        
        # Área de texto para métricas
        self.texto_metricas = tk.Text(
            text_frame,
            wrap='word',
            font=('Segoe UI', 11),
            bg=self.colores['bg_primary'],
            fg=self.colores['text_primary'],
            relief='flat',
            bd=0,
            padx=15,
            pady=15,
            spacing1=4,
            spacing2=2,
            spacing3=1,
            insertbackground=self.colores['text_primary'],
            selectbackground=self.colores['accent_green'],
            selectforeground='white',
            height=20
        )
        self.texto_metricas.pack(side='left', fill='both', expand=True)
        
        # Scrollbar
        scrollbar_metricas = tk.Scrollbar(
            text_frame,
            orient='vertical',
            command=self.texto_metricas.yview,
            bg=self.colores['bg_secondary'],
            troughcolor=self.colores['bg_input'],
            activebackground=self.colores['accent_green'],
            width=14
        )
        scrollbar_metricas.pack(side='right', fill='y', padx=(5, 0))
        self.texto_metricas.config(yscrollcommand=scrollbar_metricas.set)
        
        # Texto inicial
        self.mostrar_en_metricas("📈 MÉTRICAS DE EVALUACIÓN DEL MODELO\n" +
                                "Esta función estará disponible después de completar el análisis de sentimientos.\n\n" +
                                "📊 Métricas que se calcularán:\n" +
                                "• Exactitud (Accuracy): Porcentaje de predicciones correctas\n" +
                                "• Precisión (Precision): Exactitud de predicciones positivas\n" +
                                "• Exhaustividad (Recall): Capacidad de detectar casos positivos\n" +
                                "• F1-Score: Media armónica entre precisión y recall\n" +
                                "• Matriz de Confusión: Visualización de aciertos y errores\n" +
                                "• Métricas por clase: Análisis individual por sentimiento\n\n" +
                                "🎯 Estas métricas te ayudarán a evaluar el rendimiento\n" +
                                "de tu modelo de análisis de sentimientos.")

    def crear_footer_profesional(self, parent):
        """Crea footer profesional con progreso"""
        footer_container = tk.Frame(parent, bg=self.colores['bg_primary'])
        footer_container.pack(fill='x')
        
        # Separador superior
        sep = tk.Frame(footer_container, height=1, bg=self.colores['border'])
        sep.pack(fill='x', pady=(0, 20))
        
        # Container de progreso
        progress_container = tk.Frame(footer_container, bg=self.colores['bg_primary'])
        progress_container.pack(fill='x')
        
        # Label de estado
        status_frame = tk.Frame(progress_container, bg=self.colores['bg_primary'])
        status_frame.pack(fill='x', pady=(0, 10))
        
        tk.Label(
            status_frame,
            text="Estado del Sistema:",
            font=('Segoe UI', 10, 'bold'),
            bg=self.colores['bg_primary'],
            fg=self.colores['text_primary']
        ).pack(side='left')
        
        self.progress_label = tk.Label(
            status_frame,
            text="🚀 Sistema listo - Corrección ortográfica automática disponible",
            font=('Segoe UI', 10),
            bg=self.colores['bg_primary'],
            fg=self.colores['text_secondary']
        )
        self.progress_label.pack(side='left', padx=(10, 0))
        
        # Barra de progreso profesional
        self.progreso = ttk.Progressbar(
            progress_container,
            mode='determinate',
            style='Professional.Horizontal.TProgressbar'
        )
        self.progreso.pack(fill='x', pady=(0, 5))

    # ========== MÉTODOS DE FUNCIONALIDAD ==========

    def cargar_archivo(self):
        tipos_archivo = [
            ("Archivos de Texto", "*.txt"),
            ("Archivos CSV", "*.csv"),
            ("Archivos Excel", "*.xlsx"),
            ("Archivos Excel Legacy", "*.xls"),
            ("Archivos JSON", "*.json"),
            ("Archivos TSV", "*.tsv"),
            ("Todos los archivos", "*.*")
        ]
        
        if DOCX_DISPONIBLE:
            tipos_archivo.insert(4, ("Archivos Word", "*.docx"))
        if PDF_DISPONIBLE:
            tipos_archivo.insert(-2, ("Archivos PDF", "*.pdf"))
        
        archivo = filedialog.askopenfilename(
            title="Seleccionar archivo para análisis profesional",
            filetypes=tipos_archivo
        )
        
        if archivo:
            self.progress_label.config(text="📂 Cargando y validando archivo...")
            self.progreso.config(mode='indeterminate')
            self.progreso.start()
            self.ventana.update()
            
            def cargar():
                exito, mensaje, cantidad = self.logica.cargar_archivo(archivo)
                self.progreso.stop()
                self.progreso.config(mode='determinate', value=0)
                
                if exito:
                    self.progress_label.config(text="✅ Archivo cargado y validado correctamente")
                    self.status_badge.config(text="Cargado", bg=self.colores['accent_green'])
                    self.icon_archivo.config(fg=self.colores['accent_green'])
                    
                    nombre_archivo = os.path.basename(archivo)
                    tamaño_kb = os.path.getsize(archivo) / 1024
                    longitud_promedio = self.logica.datos['texto'].str.len().mean()
                    
                    info_text = f"📁 {nombre_archivo}\n"
                    info_text += f"📄 Total de registros: {cantidad:,}\n"
                    info_text += f"📏 Longitud promedio: {longitud_promedio:.0f} caracteres\n"
                    info_text += f"💾 Tamaño del archivo: {tamaño_kb:.1f} KB\n"
                    info_text += f"🎯 Estado: Listo para análisis\n"
                    info_text += f"✏️ Corrección ortográfica disponible"
                    
                    self.info_archivo.config(text=info_text)
                    
                    preview = "📁 VISTA PREVIA DEL CONTENIDO\n"
                    preview += "=" * 60 + "\n"
                    preview += f"Mostrando los primeros 5 registros de {cantidad:,} total:\n"
                    
                    for i, texto in enumerate(self.logica.datos['texto'].head(5), 1):
                        preview += f"📄 Registro {i}:\n"
                        preview += f"   {texto[:150]}{'...' if len(texto) > 150 else ''}\n"
                    
                    preview += f"✅ Archivo procesado exitosamente y listo para análisis.\n"
                    preview += f"✏️ La corrección ortográfica se aplicará durante la limpieza."
                    
                    self.mostrar_en_resumen(preview)
                    
                    self.btn_analizar.config(state='normal')
                    self.btn_resetear.config(state='normal')
                    
                    # Actualizar scroll después de cambios
                    self.ventana.after(100, self.actualizar_scroll)
                    
                else:
                    self.progress_label.config(text="❌ Error al procesar el archivo")
                    self.status_badge.config(text="Error", bg=self.colores['accent_red'])
                    messagebox.showerror("Error de Carga", f"No se pudo cargar el archivo:\n{mensaje}")
                
                self.ventana.update()
            
            threading.Thread(target=cargar, daemon=True).start()

    def analizar_sentimientos(self):
        self.progress_label.config(text="🔍 Ejecutando análisis avanzado de sentimientos...")
        self.progreso.config(mode='indeterminate')
        self.progreso.start()
        self.ventana.update()
        
        def analizar():
            exito, mensaje = self.logica.analizar_sentimientos()
            self.progreso.stop()
            self.progreso.config(mode='determinate', value=100)
            
            if exito:
                self.progress_label.config(text="✅ Análisis avanzado completado con éxito")
                self.analysis_badge.config(text="Completado", bg=self.colores['accent_green'])
                self.analisis_completado = True
                
                stats, resumen, datos = self.logica.generar_estadisticas()
                self.stats_label.config(text=stats)
                self.mostrar_en_resumen(resumen)
                self.mostrar_en_datos(datos)
                
                # Habilitar botones post-análisis
                self.btn_exportar.config(state='normal')
                self.btn_graficos.config(state='normal')
                self.btn_limpiar.config(state='normal')  # Habilitar limpieza robusta
                self.btn_metricas.config(state='normal')  # Habilitar métricas
                
                # Actualizar información de limpieza en el dashboard
                texto_limpieza_habilitada = "✏️ Corrección ortográfica disponible\n• Análisis completado exitosamente\n• 'me seto felis' → 'me siento feliz'\n• Conversión de emojis a texto\n• Limpieza de URLs y emails\n• Preservación del contexto emocional"
                self.limpieza_label.config(text=texto_limpieza_habilitada)
                self.cleaning_badge.config(text="Disponible", bg=self.colores['accent_blue'])
                
                # Actualizar scroll después de cambios
                self.ventana.after(100, self.actualizar_scroll)
                
            else:
                self.progress_label.config(text="❌ Error durante el análisis")
                self.analysis_badge.config(text="Error", bg=self.colores['accent_red'])
                messagebox.showerror("Error de Análisis", f"Error en el proceso:\n{mensaje}")
            
            self.progreso.config(value=0)
            self.ventana.update()
        
        threading.Thread(target=analizar, daemon=True).start()

    def calcular_metricas_evaluacion(self):
        """Calcula y muestra las métricas de evaluación del modelo"""
        if not self.analisis_completado:
            messagebox.showwarning("Análisis Requerido", 
                                 "Debes completar el análisis de sentimientos antes de calcular métricas.\n" +
                                 "Las métricas evalúan el rendimiento del modelo de análisis.")
            return
        
        self.progress_label.config(text="📊 Calculando métricas de evaluación del modelo...")
        self.progreso.config(mode='indeterminate')
        self.progreso.start()
        self.ventana.update()
        
        def calcular_metricas():
            try:
                # Calcular métricas usando el módulo importado
                metricas, reporte = calcular_metricas_modelo(
                    self.logica.datos,
                    columna_prediccion='sentimiento'
                )
                
                self.progreso.stop()
                self.progreso.config(mode='determinate', value=100)
                
                # Mostrar el reporte completo en la pestaña de métricas
                self.mostrar_en_metricas(reporte)
                
                # Actualizar interfaz
                self.progress_label.config(text="✅ Métricas de evaluación calculadas exitosamente")
                self.metricas_calculadas = True
                
                # Mostrar mensaje informativo
                exactitud = metricas['exactitud']
                f1_score = metricas['f1_macro']
                tipo_calculo = metricas.get('tipo_calculo', 'desconocido')
                
                if tipo_calculo == 'estimado':
                    messagebox.showinfo("Métricas Calculadas", 
                                      f"✅ Métricas de evaluación calculadas (estimadas)\n\n" +
                                      f"📊 Resultados principales:\n" +
                                      f"• Exactitud: {exactitud:.2%}\n" +
                                      f"• F1-Score: {f1_score:.2%}\n" +
                                      f"• Precisión: {metricas['precision_macro']:.2%}\n" +
                                      f"• Recall: {metricas['recall_macro']:.2%}\n\n" +
                                      f"⚠️ Nota: Métricas estimadas sin etiquetas ground truth\n" +
                                      f"Para métricas reales, proporcione etiquetas de referencia.")
                else:
                    messagebox.showinfo("Métricas Calculadas", 
                                      f"✅ Métricas de evaluación calculadas (reales)\n\n" +
                                      f"📊 Resultados principales:\n" +
                                      f"• Exactitud: {exactitud:.2%}\n" +
                                      f"• F1-Score: {f1_score:.2%}\n" +
                                      f"• Precisión: {metricas['precision_macro']:.2%}\n" +
                                      f"• Recall: {metricas['recall_macro']:.2%}\n\n" +
                                      f"✅ Evaluación realizada con etiquetas ground truth")
                
                # Actualizar scroll después de cambios
                self.ventana.after(100, self.actualizar_scroll)
                
            except Exception as e:
                self.progreso.stop()
                self.progreso.config(value=0)
                self.progress_label.config(text="❌ Error al calcular métricas")
                messagebox.showerror("Error en Métricas", f"Error durante el cálculo de métricas:\n{str(e)}")
            
            self.ventana.update()
        
        threading.Thread(target=calcular_metricas, daemon=True).start()

    def limpiar_datos_robusto(self):
        """Ejecuta la limpieza robusta CON CORRECCIÓN ORTOGRÁFICA usando el módulo importado"""
        if not self.analisis_completado:
            messagebox.showwarning("Análisis Requerido", 
                                 "Debes completar el análisis de sentimientos antes de usar la corrección ortográfica.\n" +
                                 "La corrección post-análisis permite preservar mejor el contexto emocional.")
            return
        
        self.progress_label.config(text="✏️ Ejecutando corrección ortográfica y limpieza robusta...")
        self.progreso.config(mode='indeterminate')
        self.progreso.start()
        self.ventana.update()
        
        def limpiar():
            # LLAMAR AL MÉTODO DEL LIMPIADOR ROBUSTO IMPORTADO
            exito, mensaje, datos_limpios, estadisticas = self.limpiador_robusto.limpiar_datos_post_analisis(self.logica.datos)
            self.progreso.stop()
            self.progreso.config(mode='determinate', value=0)
            
            if exito:
                self.progress_label.config(text="✅ Corrección ortográfica completada exitosamente")
                self.cleaning_badge.config(text="Aplicado", bg=self.colores['accent_green'])
                self.limpieza_robusta_aplicada = True
                
                # Actualizar datos en la lógica
                self.logica.datos = datos_limpios
                
                # Actualizar información de limpieza en el dashboard
                palabras_corregidas = estadisticas.get('palabras_corregidas', 0)
                textos_corregidos = estadisticas.get('textos_con_correcciones', 0)
                emojis_convertidos = estadisticas.get('emojis_convertidos', 0)
                
                limpieza_text = f"✏️ Corrección ortográfica aplicada\n"
                limpieza_text += f"📝 Palabras corregidas: {palabras_corregidas:,}\n"
                limpieza_text += f"📄 Textos corregidos: {textos_corregidos:,}\n"
                limpieza_text += f"😊 Emojis convertidos: {emojis_convertidos:,}\n"
                limpieza_text += f"🎯 Datos optimizados para exportación\n"
                limpieza_text += f"✅ Contexto emocional preservado"
                
                self.limpieza_label.config(text=limpieza_text)
                
                # GENERAR REPORTE DE LIMPIEZA USANDO EL MÉTODO DEL MÓDULO
                reporte_limpieza = self.limpiador_robusto.generar_reporte_limpieza()
                self.mostrar_en_limpieza_robusta(reporte_limpieza)
                
                # Actualizar vista previa con datos limpios
                preview_limpio = "✏️ DATOS PROCESADOS CON CORRECCIÓN ORTOGRÁFICA\n"
                preview_limpio += "=" * 60 + "\n"
                preview_limpio += f"Se corrigieron {palabras_corregidas:,} palabras en {textos_corregidos:,} textos.\n"
                preview_limpio += "Mostrando los primeros 5 registros procesados:\n\n"
                
                for i, row in enumerate(self.logica.datos.head(5).itertuples(), 1):
                    preview_limpio += f"📄 Registro {i} (corregido):\n"
                    texto_actual = row.texto
                    preview_limpio += f"   Actual: {texto_actual[:120]}{'...' if len(texto_actual) > 120 else ''}\n"
                    
                    if hasattr(row, 'texto_pre_limpieza'):
                        texto_original = row.texto_pre_limpieza
                        if texto_original != texto_actual:
                            preview_limpio += f"   Original: {texto_original[:120]}{'...' if len(texto_original) > 120 else ''}\n"
                    
                    preview_limpio += "\n"
                
                preview_limpio += f"✅ Los datos han sido optimizados con corrección ortográfica.\n"
                preview_limpio += f"🎯 Listos para exportación y análisis posterior.\n"
                preview_limpio += f"📊 Calidad de datos mejorada significativamente."
                
                self.mostrar_en_resumen(preview_limpio)
                
                # Regenerar estadísticas con datos limpios
                if hasattr(self.logica, 'resultados') and self.logica.resultados is not None:
                    stats_actualizadas, resumen_actualizado, datos_actualizados = self.logica.generar_estadisticas()
                    self.mostrar_en_datos(datos_actualizados)
                
                # Actualizar scroll después de cambios
                self.ventana.after(100, self.actualizar_scroll)
                
                # Mostrar ejemplos si los hay
                ejemplos_texto = ""
                if 'ejemplos_correcciones' in estadisticas and estadisticas['ejemplos_correcciones']:
                    ejemplos_texto = "\n\n💡 Ejemplos de correcciones realizadas:\n"
                    for i, ejemplo in enumerate(estadisticas['ejemplos_correcciones'][:3], 1):
                        ejemplos_texto += f"{i}. Original: {ejemplo['original']}\n"
                        ejemplos_texto += f"   Corregido: {ejemplo['corregido']}\n"
                
                messagebox.showinfo("Corrección Completada", 
                                  f"✅ Corrección ortográfica completada con éxito!\n\n" +
                                  f"📝 Estadísticas de corrección:\n" +
                                  f"• {palabras_corregidas:,} palabras corregidas\n" +
                                  f"• {textos_corregidos:,} textos modificados\n" +
                                  f"• {emojis_convertidos:,} emojis convertidos\n" +
                                  f"• {estadisticas.get('urls_removidas', 0):,} URLs eliminadas\n" +
                                  f"• {estadisticas.get('emails_removidos', 0):,} emails eliminados\n\n" +
                                  f"🎯 Los datos están ahora corregidos y listos para exportación.\n" +
                                  f"📈 La calidad de los datos ha mejorado significativamente." + ejemplos_texto)
            else:
                self.progress_label.config(text="❌ Error en la corrección ortográfica")
                messagebox.showerror("Error de Corrección", f"Error durante la corrección ortográfica:\n{mensaje}")
            
            self.ventana.update()
        
        threading.Thread(target=limpiar, daemon=True).start()

    def mostrar_graficos(self):
        self.progress_label.config(text="📊 Generando visualizaciones avanzadas...")
        self.progreso.config(mode='indeterminate')
        self.progreso.start()
        self.ventana.update()
        
        def mostrar():
            self.logica.mostrar_graficos()
            self.progreso.stop()
            self.progreso.config(value=0)
            self.progress_label.config(text="✅ Gráficos generados correctamente")
            self.ventana.update()
        
        threading.Thread(target=mostrar, daemon=True).start()

    def exportar_resultados(self):
        archivo_salida = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[
                ("Excel Profesional", "*.xlsx"),
                ("CSV Datos", "*.csv"),
                ("JSON Estructurado", "*.json"),
                ("Reporte Texto", "*.txt")
            ],
            title="Exportar Resultados Corregidos del Análisis"
        )
        
        if archivo_salida:
            self.progress_label.config(text="💾 Exportando resultados con correcciones...")
            self.progreso.config(mode='indeterminate')
            self.progreso.start()
            
            def exportar():
                exito, mensaje = self.logica.exportar_resultados(archivo_salida)
                self.progreso.stop()
                self.progreso.config(value=0)
                
                if exito:
                    self.progress_label.config(text="✅ Exportación completada exitosamente")
                    
                    detalles_exportacion = f"✅ Resultados exportados correctamente:\n"
                    detalles_exportacion += f"📁 {os.path.basename(archivo_salida)}\n"
                    detalles_exportacion += f"📂 {os.path.dirname(archivo_salida)}\n\n"
                    detalles_exportacion += f"📊 El archivo incluye:\n"
                    detalles_exportacion += f"• Resultados completos del análisis de sentimientos\n"
                    detalles_exportacion += f"• Métricas avanzadas y estadísticas detalladas\n"
                    
                    if self.limpieza_robusta_aplicada:
                        detalles_exportacion += f"• Textos con corrección ortográfica aplicada\n"
                        detalles_exportacion += f"• Comparación antes/después de corrección\n"
                        detalles_exportacion += f"• Emojis convertidos a texto descriptivo\n"
                    
                    if self.metricas_calculadas:
                        detalles_exportacion += f"• Métricas de evaluación del modelo\n"
                        detalles_exportacion += f"• Reporte detallado de rendimiento\n"
                    
                    detalles_exportacion += f"• Información de procesamiento y calidad\n"
                    detalles_exportacion += f"• Datos categorizados por tipo de sentimiento\n"
                    detalles_exportacion += f"• Textos limpios y listos para uso profesional"
                    
                    messagebox.showinfo("Exportación Exitosa", detalles_exportacion)
                else:
                    self.progress_label.config(text="❌ Error en exportación")
                    messagebox.showerror("Error de Exportación", f"Error al exportar:\n{mensaje}")
                
                self.ventana.update()
            
            threading.Thread(target=exportar, daemon=True).start()

    def resetear_analisis(self):
        """Resetea completamente el análisis"""
        respuesta = messagebox.askyesno("Confirmar Reset", 
                                      "¿Estás seguro de que quieres resetear todo el análisis?\n\n" +
                                      "Se perderán:\n" +
                                      "• Todos los datos cargados\n" +
                                      "• Resultados del análisis\n" +
                                      "• Correcciones ortográficas aplicadas\n" +
                                      "• Métricas calculadas\n" +
                                      "• Configuraciones actuales")
        
        if respuesta:
            # Resetear datos
            self.logica.datos = None
            self.logica.datos_originales = None
            self.logica.resultados = None
            if hasattr(self.logica, 'correcciones_realizadas'):
                self.logica.correcciones_realizadas = {}
            
            # Resetear estados
            self.correcciones_aplicadas = False
            self.analisis_completado = False
            self.limpieza_robusta_aplicada = False
            self.metricas_calculadas = False
            
            # RESETEAR LIMPIADOR CREANDO UNA NUEVA INSTANCIA DEL MÓDULO
            self.limpiador_robusto = LimpiadorDatosRobusto()
            # RESETEAR CALCULADOR DE MÉTRICAS
            self.calculador_metricas = CalculadorMetricas()
            
            # Resetear interfaz - MANTIENE TEXTO ORIGINAL DEL DASHBOARD
            self.info_archivo.config(text="📁 No hay archivo seleccionado\n• Selecciona un archivo de texto, CSV, Excel o Word\n• Formatos soportados: TXT, CSV, XLSX, DOCX, PDF\n• Tamaño máximo recomendado: 50MB\n• Corrección ortográfica automática disponible")
            self.stats_label.config(text="📊 Ejecuta el análisis para ver estadísticas\n• Distribución de sentimientos\n• Puntuaciones de confianza\n• Métricas de intensidad emocional\n• Análisis de correlaciones\n• Insights detallados")
            
            # MANTIENE TEXTO ORIGINAL DE CORRECCIÓN ORTOGRÁFICA
            texto_limpieza_original = "✏️ Corrección automática de ortografía\n• Disponible tras completar análisis\n• 'me seto felis' → 'me siento feliz'\n• Conversión de emojis a texto\n• Limpieza de URLs y emails\n• Preservación del contexto emocional"
            self.limpieza_label.config(text=texto_limpieza_original)
            
            # Resetear badges
            self.status_badge.config(text="Esperando", bg=self.colores['text_muted'])
            self.analysis_badge.config(text="Pendiente", bg=self.colores['text_muted'])
            self.cleaning_badge.config(text="Pendiente", bg=self.colores['text_muted'])
            
            # Resetear iconos
            self.icon_archivo.config(fg=self.colores['text_secondary'])
            self.icon_limpieza.config(text="✏️", fg=self.colores['accent_purple'])
            
            # Limpiar áreas de texto
            self.mostrar_en_resumen("📋 RESUMEN EJECUTIVO\nAquí se mostrará un resumen completo del análisis una vez que cargues un archivo y ejecutes el análisis de sentimientos con corrección ortográfica automática.")
            self.mostrar_en_datos("📊 DATOS DETALLADOS\nAquí se mostrarán los datos detallados del análisis con todas las métricas calculadas y correcciones ortográficas aplicadas.")
            self.mostrar_en_limpieza_robusta("✏️ CORRECCIÓN ORTOGRÁFICA AUTOMÁTICA\nEsta función estará disponible después de completar el análisis de sentimientos.\n📝 La corrección ortográfica incluye:\n• Corrección de palabras mal escritas: 'me seto felis' → 'me siento feliz'\n• Normalización de texto: 'estoi' → 'estoy'\n• Conversión de emojis a texto descriptivo\n• Eliminación de URLs y enlaces\n• Limpieza de direcciones de email\n• Normalización de caracteres especiales\n• Preservación del contexto emocional")
            self.mostrar_en_metricas("📈 MÉTRICAS DE EVALUACIÓN DEL MODELO\nEsta función estará disponible después de completar el análisis de sentimientos.\n\n📊 Métricas que se calcularán:\n• Exactitud (Accuracy): Porcentaje de predicciones correctas\n• Precisión (Precision): Exactitud de predicciones positivas\n• Exhaustividad (Recall): Capacidad de detectar casos positivos\n• F1-Score: Media armónica entre precisión y recall\n• Matriz de Confusión: Visualización de aciertos y errores\n• Métricas por clase: Análisis individual por sentimiento\n\n🎯 Estas métricas te ayudarán a evaluar el rendimiento\nde tu modelo de análisis de sentimientos.")
            
            # Resetear botones
            self.btn_analizar.config(state='disabled')
            self.btn_limpiar.config(state='disabled')
            self.btn_metricas.config(state='disabled')
            self.btn_exportar.config(state='disabled')
            self.btn_graficos.config(state='disabled')
            self.btn_resetear.config(state='disabled')
            
            # Resetear barra de progreso y mensaje
            self.progreso.config(value=0)
            self.progress_label.config(text="🚀 Sistema reseteado - Corrección ortográfica automática disponible")
            
            # Actualizar scroll después de reset
            self.ventana.after(100, self.actualizar_scroll)
            
            messagebox.showinfo("Reset Completado", 
                              "✅ El sistema ha sido reseteado exitosamente.\n\n" +
                              "🔄 Estado actual:\n" +
                              "• Todas las configuraciones restauradas\n" +
                              "• Memoria liberada completamente\n" +
                              "• Sistema listo para nuevo análisis\n" +
                              "• Corrección ortográfica disponible\n" +
                              "• Evaluación de métricas disponible\n\n" +
                              "📁 Puedes cargar un nuevo archivo para comenzar.")

    def actualizar_scroll(self):
        """Método para actualizar el scroll manualmente"""
        self.canvas.update_idletasks()
        self.canvas.configure(scrollregion=self.canvas.bbox("all"))
        
        # Configurar el ancho del canvas window
        canvas_width = self.canvas.winfo_width()
        if canvas_width > 1:
            self.canvas.itemconfig(self.canvas_window, width=canvas_width)

    def mostrar_en_resumen(self, texto):
        """Muestra texto en la pestaña de resumen y actualiza el scroll"""
        self.texto_resumen.config(state='normal')
        self.texto_resumen.delete(1.0, tk.END)
        self.texto_resumen.insert(1.0, texto)
        self.texto_resumen.config(state='disabled')
        self.ventana.after(50, self.actualizar_scroll)

    def mostrar_en_datos(self, texto):
        """Muestra texto en la pestaña de datos y actualiza el scroll"""
        self.texto_datos.config(state='normal')
        self.texto_datos.delete(1.0, tk.END)
        self.texto_datos.insert(1.0, texto)
        self.texto_datos.config(state='disabled')
        self.ventana.after(50, self.actualizar_scroll)

    def mostrar_en_limpieza_robusta(self, texto):
        """Muestra texto en la pestaña de limpieza robusta y actualiza el scroll"""
        self.texto_limpieza_robusta.config(state='normal')
        self.texto_limpieza_robusta.delete(1.0, tk.END)
        self.texto_limpieza_robusta.insert(1.0, texto)
        self.texto_limpieza_robusta.config(state='disabled')
        self.ventana.after(50, self.actualizar_scroll)

    def mostrar_en_metricas(self, texto):
        """Muestra texto en la pestaña de métricas y actualiza el scroll"""
        self.texto_metricas.config(state='normal')
        self.texto_metricas.delete(1.0, tk.END)
        self.texto_metricas.insert(1.0, texto)
        self.texto_metricas.config(state='disabled')
        self.ventana.after(50, self.actualizar_scroll)

    def ejecutar(self):
        """Ejecuta la aplicación principal"""
        # Configuración final del scroll antes de mostrar
        self.ventana.after(200, self.actualizar_scroll)
        self.ventana.mainloop()


if __name__ == "__main__":
    print("✨ ANALIZADOR DE SENTIMIENTOS PROFESIONAL v3.0 CON CORRECCIÓN ORTOGRÁFICA Y MÉTRICAS")
    print("=" * 80)
    print("🚀 Inicializando interfaz profesional con corrección ortográfica y evaluación de métricas...")
    print("✏️ Sistema de corrección ortográfica automática integrado")
    print("📊 Módulo de evaluación de métricas integrado")
    
    # NUEVO: Mostrar información de VADER y Naive Bayes
    if VADER_DISPONIBLE:
        print("⚡ VADER Sentiment Analysis disponible")
    else:
        print("⚠️ VADER no disponible - Instala con: pip install nltk")
    
    if SKLEARN_DISPONIBLE:
        print("📊 Naive Bayes (Bernoulli) disponible")
    else:
        print("⚠️ Naive Bayes no disponible - Instala con: pip install scikit-learn")
    
    if not verificar_dependencias():
        print("\n❌ Error: Dependencias faltantes.")
        print("   Instala los paquetes requeridos antes de continuar.")
        exit(1)
    
    try:
        app = AnalizadorSentimientosGUI()
        print("✅ Interfaz cargada correctamente")
        
        if SPELLCHECKER_DISPONIBLE:
            print("✏️ Sistema de corrección ortográfica disponible")
            print("   📝 Ejemplos: 'me seto felis' → 'me siento feliz'")
        else:
            print("⚠️ Sistema de corrección ortográfica no disponible")
            print("   Instala con: pip install pyspellchecker")
        
        print("📊 Módulo de métricas de evaluación integrado")
        print("   🎯 Métricas: Exactitud, Precisión, Recall, F1-Score")
        print("   📈 Evaluación completa del modelo")
        
        print("🧹 Sistema de limpieza robusta con corrección ortográfica activado")
        print("🎯 Iniciando aplicación avanzada...")
        app.ejecutar()
        
    except Exception as e:
        print(f"❌ Error crítico al iniciar la aplicación: {e}")
        print("💡 Verifica que todas las dependencias estén instaladas correctamente.")
        print("🔥 Dependencias requeridas: pandas numpy textblob matplotlib seaborn openpyxl xlrd")
        print("🔥 Dependencias opcionales: python-docx PyPDF2 wordcloud nltk pyspellchecker")